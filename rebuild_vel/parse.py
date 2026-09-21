"""Parse merged procedure-data xlsx exports into structured measurement records.

The platform's ``/flow/procedureDataMergeExport`` endpoint returns one workbook
per ``(station, device, measureTimeList)`` chunk with five sheets:

* 水位流量数据表          - one row per measurement (levels, areas, flows)
* 测速线流速数据表        - final per-line surface velocity (the rebuild target)
* STIV原始流速数据表      - raw per-video-segment velocities + confidence
* 光流轨迹法原始流速数据表 - raw optical-flow trajectory rows (mode 3 devices)
* 光流法原始流速数据表      - raw dense optical-flow rows (mode 3 devices)

All numeric cells arrive as strings; everything is parsed defensively.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Optional

import numpy as np

SHEET_LEVEL_FLOW = "水位流量数据表"
SHEET_LINES = "测速线流速数据表"
SHEET_STIV_RAW = "STIV原始流速数据表"
SHEET_OF_TRAJ_RAW = "光流轨迹法原始流速数据表"
SHEET_OF_RAW = "光流法原始流速数据表"

_TIME_RE = re.compile(r"(\d{4})[-/](\d{2})[-/](\d{2})\s+(\d{2}):(\d{2})(?::(\d{2}))?")


def normalize_time(value: Any) -> str:
    """Normalize ``2025/09/30 23:30`` / ``2025-09-30 23:30:00`` to ``YYYY-MM-DD HH:MM:SS``."""
    text = str(value or "").strip()
    match = _TIME_RE.search(text)
    if not match:
        return text
    year, month, day, hour, minute, second = match.groups()
    return f"{year}-{month}-{day} {hour}:{minute}:{second or '00'}"


def _float(value: Any, default: float = np.nan) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in ("", "--", "-", "null", "None"):
        return default
    try:
        return float(text)
    except ValueError:
        return default


def _sheet_rows(workbook: Any, name: str) -> tuple[list[str], list[tuple[Any, ...]]]:
    if name not in workbook.sheetnames:
        return [], []
    sheet = workbook[name]
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        return [], []
    header = [str(cell or "").strip() for cell in rows[0]]
    body = [row for row in rows[1:] if any(cell not in (None, "") for cell in row)]
    return header, body


@dataclass
class ParsedExport:
    """One merged xlsx parsed into per-measurement records."""

    measurements: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_merged_export(payload: bytes) -> ParsedExport:
    import openpyxl  # deferred: heavy import, only needed when parsing

    result = ParsedExport()
    if not payload or payload[:2] != b"PK":
        result.errors.append("payload is not an xlsx workbook")
        return result
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        result.errors.append(f"failed to open workbook: {exc}")
        return result

    level_header, level_rows = _sheet_rows(workbook, SHEET_LEVEL_FLOW)
    line_header, line_rows = _sheet_rows(workbook, SHEET_LINES)
    stiv_header, stiv_rows = _sheet_rows(workbook, SHEET_STIV_RAW)
    traj_header, traj_rows = _sheet_rows(workbook, SHEET_OF_TRAJ_RAW)
    of_header, of_rows = _sheet_rows(workbook, SHEET_OF_RAW)

    measurements: dict[tuple[str, str, str], dict[str, Any]] = {}

    def record(station: str, device: str, time_value: Any) -> dict[str, Any]:
        key = (station, device, normalize_time(time_value))
        if key not in measurements:
            measurements[key] = {
                "station": key[0],
                "device": key[1],
                "time": key[2],
                "station_name": "",
                "water_level": np.nan,
                "section_area": np.nan,
                "water_width": np.nan,
                "surface_avg_velocity": np.nan,
                "max_depth": np.nan,
                "aver_depth": np.nan,
                "max_surface_velocity": np.nan,
                "original_value_prop": np.nan,
                "version": "",
                "lines": [],
                "raw_stiv": [],
                "raw_of_traj": [],
                "raw_of": [],
            }
        return measurements[key]

    # -- level/flow sheet --------------------------------------------------
    if level_rows:
        idx = {name: level_header.index(name) for name in level_header}
        for row in level_rows:
            station = str(row[idx.get("站名", 0)] or "").strip()
            device = str(row[idx.get("设备码", 1)] or "").strip()
            rec = record(station, device, row[idx.get("测流时间", 2)])
            rec["station_name"] = station
            rec["water_level"] = _float(row[idx.get("水位", 3)])
            rec["section_area"] = _float(row[idx.get("断面面积", 4)])
            rec["water_width"] = _float(row[idx.get("河宽", 5)])
            rec["surface_avg_velocity"] = _float(row[idx.get("表面平均流速", 7)])
            rec["aver_depth"] = _float(row[idx.get("平均水深", 15)])
            rec["max_depth"] = _float(row[idx.get("最大水深", 16)])
            rec["max_surface_velocity"] = _float(row[idx.get("最大表面流速", 17)])
            rec["original_value_prop"] = _float(row[idx.get("流速原始真值占比", 19)])

    # -- final speed lines (the rebuild target) ----------------------------
    if line_rows:
        idx = {name: line_header.index(name) for name in line_header}

        def cell(row: tuple[Any, ...], name: str, fallback: int) -> Any:
            """Read a column by header name; short rows degrade to empty
            instead of raising IndexError (some exports drop tail columns)."""
            i = idx.get(name, fallback)
            return row[i] if i < len(row) else ""

        for row in line_rows:
            station = str(cell(row, "站名", 0) or "").strip()
            device = str(cell(row, "设备码", 1) or "").strip()
            rec = record(station, device, cell(row, "测流时间", 2))
            coords = str(cell(row, "测速线坐标", 19) or "").strip()
            coord_parts = [_float(part) for part in coords.split("-")] if coords else []
            rec["lines"].append({
                "line_num": _float(cell(row, "测速线序号", 8), -1.0),
                "x": _float(cell(row, "测速线起点距", 9)),
                "bed_elevation": _float(cell(row, "测速线河底高程", 10)),
                "area": _float(cell(row, "测速线面积", 11)),
                "depth": _float(cell(row, "测速线水深", 12)),
                "v_surface": _float(cell(row, "测速线流速", 13)),
                "flow": _float(cell(row, "测速线流量", 14)),
                "confidence": _float(cell(row, "测速线平均置信度", 15)),
                "angle": _float(cell(row, "测速线平均流向夹角", 16)),
                "is_algo": _float(cell(row, "线流速类型1原始值0插值", 17), 0.0) == 1.0,
                "version": str(cell(row, "算法版本", 18) or "").strip(),
                "coords": coord_parts,
            })
            if not rec["version"]:
                rec["version"] = str(cell(row, "算法版本", 18) or "").strip()

    def _raw_row(idx: dict[str, int], row: tuple[Any, ...], extra: dict[str, Any]) -> dict[str, Any]:
        return {
            "work_point": _float(row[idx["工作点编号"]]) if "工作点编号" in idx else np.nan,
            "t_start": _float(row[idx["起始时间"]]) if "起始时间" in idx else np.nan,
            "t_end": _float(row[idx["终止时间"]]) if "终止时间" in idx else np.nan,
            "line_num": _float(row[idx["测速线序号"]], -1.0),
            "x": _float(row[idx["起点距"]]),
            "raw_v": _float(row[idx["原始流速"]]),
            "depth": _float(row[idx["水深"]]) if "水深" in idx else np.nan,
            "confidence": _float(row[idx["置信度"]]) if "置信度" in idx else np.nan,
            "angle": _float(row[idx["流向夹角"]]) if "流向夹角" in idx else np.nan,
            "version": str(row[idx["算法版本"]]).strip() if "算法版本" in idx else "",
            **extra,
        }

    # -- raw STIV rows ------------------------------------------------------
    if stiv_rows:
        idx = {name: stiv_header.index(name) for name in stiv_header}
        for row in stiv_rows:
            station = str(row[idx.get("站名", 0)] or "").strip()
            device = str(row[idx.get("设备码", 1)] or "").strip()
            rec = record(station, device, row[idx.get("测流时间", 2)])
            rec["raw_stiv"].append(_raw_row(idx, row, {
                "pixel_length": _float(row[idx["测速线像素长度"]]) if "测速线像素长度" in idx else np.nan,
                "physical_length": _float(row[idx["测速线物理长度"]]) if "测速线物理长度" in idx else np.nan,
                "line_angle": _float(row[idx["测速线识别角度"]]) if "测速线识别角度" in idx else np.nan,
                "model_angle": _float(row[idx["模型识别角度"]]) if "模型识别角度" in idx else np.nan,
                "video_segment_id": _float(row[idx["视频分段ID"]]) if "视频分段ID" in idx else np.nan,
            }))

    # -- raw optical-flow trajectory rows -----------------------------------
    if traj_rows:
        idx = {name: traj_header.index(name) for name in traj_header}
        for row in traj_rows:
            station = str(row[idx.get("站名", 0)] or "").strip()
            device = str(row[idx.get("设备码", 1)] or "").strip()
            rec = record(station, device, row[idx.get("测流时间", 2)])
            rec["raw_of_traj"].append(_raw_row(idx, row, {
                "vertical_v": _float(row[idx["垂向流速"]]) if "垂向流速" in idx else np.nan,
            }))

    # -- raw dense optical-flow rows ----------------------------------------
    if of_rows:
        idx = {name: of_header.index(name) for name in of_header}
        for row in of_rows:
            station = str(row[idx.get("站名", 0)] or "").strip()
            device = str(row[idx.get("设备码", 1)] or "").strip()
            rec = record(station, device, row[idx.get("测流时间", 2)])
            rec["raw_of"].append(_raw_row(idx, row, {
                "region_mean": _float(row[idx["区域均值"]]) if "区域均值" in idx else np.nan,
                "std": _float(row[idx["标准差"]]) if "标准差" in idx else np.nan,
                "max_v": _float(row[idx["最大流速"]]) if "最大流速" in idx else np.nan,
                "min_v": _float(row[idx["最小流速"]]) if "最小流速" in idx else np.nan,
                "direction_ratio": _float(row[idx["方向分布比"]]) if "方向分布比" in idx else np.nan,
                "pixel_scale": _float(row[idx["像素尺度"]]) if "像素尺度" in idx else np.nan,
            }))

    result.measurements = list(measurements.values())
    return result
