"""Survey which raw-velocity sheets each device/mode actually populates.

The procedure workbook carries three raw sources:

* STIV原始流速数据表        -> parsed as ``raw_stiv``   (currently the only one used)
* 光流轨迹法原始流速数据表  -> parsed as ``raw_of_traj``
* 光流法原始流速数据表      -> parsed as ``raw_of``

Every device measurement has at least one of them populated; some devices have
only the optical-flow ones. This tool samples devices across calculation modes
and reports the row counts per sheet plus the column structure, so the input
builder can be fixed against facts rather than assumptions.

Usage:
    python tools/probe_sheets.py --output data/sheet_survey.json [--devices 40] [--per-device 2]
"""

from __future__ import annotations

import argparse
import io
import json
import logging
import random
import sys
import threading
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from rebuild_vel.accounts import AccountPool
from rebuild_vel.catalog import load_catalog
from rebuild_vel.client import FlowClient, PermanentAPIError, RemoteAPIError

logger = logging.getLogger("sheets")

STATION_INFO = r"D:\chengs\vmodel\v_data\read_data\station_info_view.csv"
FLOW_CONFIG = r"D:\chengs\vmodel\v_data\read_data\calibration_fore_config_flow_view.csv"

RAW_SHEETS = {
    "stiv": "STIV原始流速数据表",
    "of_traj": "光流轨迹法原始流速数据表",
    "of": "光流法原始流速数据表",
}
KEEP_HEADERS = ["站名", "设备码", "测流时间", "水位", "工作点编号", "起始时间", "终止时间",
                "测速线序号", "起点距", "原始流速", "水深", "置信度", "流向夹角",
                "视频分段ID", "垂向流速", "区域均值"]


def sheet_stats(workbook, name: str) -> dict:
    if name not in workbook.sheetnames:
        return {"rows": -1, "header": [], "sample": []}
    rows = list(workbook[name].iter_rows(values_only=True))
    header = [str(x or "").strip() for x in rows[0]] if rows else []
    body = [r for r in rows[1:] if any(c not in (None, "") for c in r)]
    return {
        "rows": len(body),
        "header": header,
        "sample": [[str(c)[:24] for c in r] for r in body[:2]],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--devices", type=int, default=40)
    parser.add_argument("--per-device", type=int, default=2)
    parser.add_argument("--begin-time", default="2026-01-01 00:00:00.000")
    parser.add_argument("--end-time", default="2026-08-31 23:59:59.999")
    parser.add_argument("--day-only", action="store_true", default=True)
    parser.add_argument("--all-hours", dest="day_only", action="store_false")
    parser.add_argument("--seed", type=int, default=11)
    parser.add_argument("--workers", type=int, default=5)
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

    catalog = load_catalog(STATION_INFO, FLOW_CONFIG)
    by_mode: dict[str, list] = {}
    for entry in catalog:
        by_mode.setdefault(entry.calculation_mode, []).append(entry)
    rng = random.Random(args.seed)
    picked = []
    per_mode = max(1, args.devices // max(1, len(by_mode)))
    for mode, entries in sorted(by_mode.items()):
        picked.extend(rng.sample(entries, min(per_mode, len(entries))))
    logger.info("sampling %d devices across modes %s", len(picked), sorted(by_mode))

    local = threading.local()

    def client_for() -> FlowClient:
        if not hasattr(local, "client"):
            local.client = FlowClient(AccountPool("accounts.json"), timeout=90, request_delay=0.05)
        return local.client

    def survey(entry) -> dict:
        client = client_for()
        request = {
            "beginTime": args.begin_time, "endTime": args.end_time,
            "stationCode": entry.station_code, "deviceCode": entry.device_code,
            "measureResult": 1,
        }
        if args.day_only:
            request["nightMark"] = 0
        out = {"station": entry.station_code, "device": entry.device_code,
               "mode": entry.calculation_mode, "station_name": entry.station_name,
               "measurements": []}
        try:
            payload = client.post("/flow/originalDataFilterPage", {
                "count": args.per_device, "page": 1, "request": request})
        except (RemoteAPIError, PermanentAPIError) as exc:
            out["error"] = f"list: {exc}"
            return out
        rows = payload.get("data") or []
        for row in rows:
            when = str(row.get("measureTime") or "")
            rec = {"time": when, "nightMark": row.get("nightMark"),
                   "reliability": row.get("reliability"), "dataType": row.get("dataType"),
                   "originalValueProp": row.get("originalValueProp")}
            try:
                body = client.get_bytes("/flow/procedureDataMergeExport", {
                    "stationCode": entry.station_code, "deviceCode": entry.device_code,
                    "measureTimeList": [when]})
            except (RemoteAPIError, PermanentAPIError) as exc:
                rec["error"] = f"export: {exc}"
                out["measurements"].append(rec)
                continue
            if not body:
                rec["error"] = "empty export"
                out["measurements"].append(rec)
                continue
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(body))
            rec["sheets"] = {key: sheet_stats(workbook, name)["rows"]
                             for key, name in RAW_SHEETS.items()}
            rec["headers"] = {key: sheet_stats(workbook, name)["header"]
                              for key, name in RAW_SHEETS.items()}
            rec["samples"] = {key: sheet_stats(workbook, name)["sample"]
                              for key, name in RAW_SHEETS.items()}
            out["measurements"].append(rec)
        return out

    results = []
    started = time.time()
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [pool.submit(survey, e) for e in picked]
        for done, future in enumerate(as_completed(futures), 1):
            try:
                results.append(future.result())
            except Exception as exc:  # noqa: BLE001
                logger.warning("survey crashed: %s", exc)
            if done % 10 == 0:
                logger.info("surveyed %d/%d (%.0fs)", done, len(picked), time.time() - started)

    # -- summarise ---------------------------------------------------------
    combos = Counter()
    by_mode_combo = {}
    for device in results:
        keys = Counter()
        for rec in device["measurements"]:
            if "sheets" not in rec:
                keys["EXPORT_FAILED"] += 1
                continue
            populated = "+".join(k for k, v in rec["sheets"].items() if v > 0) or "ALL_EMPTY"
            keys[populated] += 1
        combos.update(keys)
        for combo in keys:
            by_mode_combo.setdefault(device["mode"], Counter())[combo] += 1

    logger.info("raw-sheet combinations seen: %s", dict(combos))
    for mode in sorted(by_mode_combo):
        logger.info("  mode %s: %s", mode, dict(by_mode_combo[mode]))

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(json.dumps({
        "day_only": args.day_only,
        "combinations": dict(combos),
        "by_mode": {m: dict(c) for m, c in by_mode_combo.items()},
        "devices": results,
        "seconds": time.time() - started,
    }, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    logger.info("wrote %s", args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
